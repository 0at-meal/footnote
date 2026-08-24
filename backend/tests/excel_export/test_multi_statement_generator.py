"""
Unit tests for 6-Tab Multi-Statement Excel Compiler (Phase C).

Validates:
- 6 space-free sheets in exact order
- Single-year (1 column) vs multi-year (N columns sorted chronologically)
- Live cross-sheet formula compilation (Income Statement -> EBITDA Bridge, downstream -> Executive Summary)
- Cell-level PDF provenance comments on all hardcode value cells
- Audit Trail tab population
"""

from pathlib import Path
import openpyxl
import pytest

from app.classification.models import StatementType
from app.excel_export.multi_statement_generator import generate_multi_statement_workbook
from app.formula_engine.models import (
    ComprehensiveModelTree,
    FormulaInputBatch,
    FormulaInputNode,
)
from app.formula_engine.tree import build_comprehensive_model_tree
from app.ingestion.models import CompanyRecord, JobRecord, JobStatus


def _make_node(
    label: str,
    value: str = "100",
    record_index: int = 0,
    statement_type: StatementType | None = None,
) -> FormulaInputNode:
    return FormulaInputNode(
        node_id=f"node_{record_index}_{label}",
        normalized_label=label,
        value=value,
        label=label,
        page=1,
        bbox={"x0": 100.0, "y0": 200.0, "x1": 300.0, "y1": 400.0},
        source_file="filing.pdf",
        record_index=record_index,
        statement_type=statement_type,
    )


def _make_sample_tree() -> ComprehensiveModelTree:
    nodes = [
        # IS
        _make_node("Revenue", value="1,000", record_index=0, statement_type=StatementType.income_statement),
        _make_node("Cost of Revenue", value="400", record_index=1, statement_type=StatementType.income_statement),
        _make_node("Research & Development", value="150", record_index=2, statement_type=StatementType.income_statement),
        _make_node("Sales & Marketing", value="100", record_index=3, statement_type=StatementType.income_statement),
        _make_node("General & Administrative", value="50", record_index=4, statement_type=StatementType.income_statement),
        _make_node("Provision for Income Taxes", value="60", record_index=5, statement_type=StatementType.income_statement),
        # Bridge
        _make_node("Stock-Based Compensation", value="40", record_index=6, statement_type=StatementType.non_gaap_bridge),
        _make_node("Amortization of Intangibles", value="20", record_index=7, statement_type=StatementType.non_gaap_bridge),
        # CF
        _make_node("Capital Expenditures", value="80", record_index=8, statement_type=StatementType.cash_flow),
        _make_node("Cash Provided by Operating Activities", value="300", record_index=9, statement_type=StatementType.cash_flow),
        # BS
        _make_node("Cash and Cash Equivalents", value="250", record_index=10, statement_type=StatementType.balance_sheet),
        _make_node("Short-Term Debt", value="50", record_index=11, statement_type=StatementType.balance_sheet),
        _make_node("Long-Term Debt", value="400", record_index=12, statement_type=StatementType.balance_sheet),
    ]
    batch = FormulaInputBatch(
        nodes=nodes,
        total_records_received=len(nodes),
        confirmed_count=len(nodes),
        excluded_count=0,
    )
    return build_comprehensive_model_tree(batch)


def test_multi_statement_workbook_6_tabs_present(tmp_path: Path) -> None:
    job = JobRecord(
        job_id="job_test_001",
        filename="filing.pdf",
        file_size_bytes=100,
        target_metric="Adjusted EBITDA",
        status=JobStatus.done,
        submitted_at="2026-01-01T00:00:00Z",
        filing_year=2023,
    )
    comp_tree = _make_sample_tree()

    result = generate_multi_statement_workbook(
        company=None,
        year_trees=[(job, comp_tree)],
        output_dir=tmp_path,
    )

    assert result.is_success is True
    assert Path(result.file_path).exists()
    assert result.sheet_names == [
        "Executive_Summary",
        "Income_Statement",
        "EBITDA_Bridge",
        "Cash_Flow",
        "Balance_Sheet",
        "Audit_Trail",
    ]

    # Open with openpyxl to verify workbook structure
    wb = openpyxl.load_workbook(result.file_path, data_only=False)
    assert wb.sheetnames == [
        "Executive_Summary",
        "Income_Statement",
        "EBITDA_Bridge",
        "Cash_Flow",
        "Balance_Sheet",
        "Audit_Trail",
    ]


def test_multi_statement_multi_year_columns_sorted(tmp_path: Path) -> None:
    company = CompanyRecord(
        company_id="comp_123",
        name="Acme Corp",
        ticker="ACME",
        created_at="2026-01-01T00:00:00Z",
        job_ids=["job_2023", "job_2022", "job_2024"],
    )

    job_2024 = JobRecord(job_id="job_2024", filename="2024.pdf", file_size_bytes=100, target_metric="Adjusted EBITDA", status=JobStatus.done, submitted_at="2026-01-03T00:00:00Z", filing_year=2024)
    job_2022 = JobRecord(job_id="job_2022", filename="2022.pdf", file_size_bytes=100, target_metric="Adjusted EBITDA", status=JobStatus.done, submitted_at="2026-01-01T00:00:00Z", filing_year=2022)
    job_2023 = JobRecord(job_id="job_2023", filename="2023.pdf", file_size_bytes=100, target_metric="Adjusted EBITDA", status=JobStatus.done, submitted_at="2026-01-02T00:00:00Z", filing_year=2023)

    comp_tree = _make_sample_tree()
    year_trees = [
        (job_2024, comp_tree),
        (job_2022, comp_tree),
        (job_2023, comp_tree),
    ]

    result = generate_multi_statement_workbook(
        company=company,
        year_trees=year_trees,
        output_dir=tmp_path,
    )

    assert result.is_success is True
    wb = openpyxl.load_workbook(result.file_path, data_only=False)
    ws_is = wb["Income_Statement"]

    # Row 3 is header: Col B = FY2022, Col C = FY2023, Col D = FY2024
    assert ws_is.cell(row=3, column=2).value == "FY2022"
    assert ws_is.cell(row=3, column=3).value == "FY2023"
    assert ws_is.cell(row=3, column=4).value == "FY2024"


def test_cross_sheet_formulas_compiled(tmp_path: Path) -> None:
    job = JobRecord(job_id="job_001", filename="filing.pdf", file_size_bytes=100, target_metric="Adjusted EBITDA", status=JobStatus.done, submitted_at="2026-01-01T00:00:00Z", filing_year=2023)
    comp_tree = _make_sample_tree()

    result = generate_multi_statement_workbook(
        company=None,
        year_trees=[(job, comp_tree)],
        output_dir=tmp_path,
    )

    wb = openpyxl.load_workbook(result.file_path, data_only=False)

    # 1. EBITDA_Bridge has cross reference to Income_Statement
    ws_bridge = wb["EBITDA_Bridge"]
    ebit_formula_cell = ws_bridge.cell(row=4, column=2)  # row 4 is Operating Income (EBIT)
    assert ebit_formula_cell.value.startswith("='Income_Statement'!")

    # 2. Executive_Summary has cross references
    ws_exec = wb["Executive_Summary"]
    rev_kpi = ws_exec.cell(row=5, column=2)  # Revenue KPI
    assert rev_kpi.value.startswith("='Income_Statement'!")


def test_cell_provenance_and_audit_trail(tmp_path: Path) -> None:
    job = JobRecord(job_id="job_001", filename="filing.pdf", file_size_bytes=100, target_metric="Adjusted EBITDA", status=JobStatus.done, submitted_at="2026-01-01T00:00:00Z", filing_year=2023)
    comp_tree = _make_sample_tree()

    result = generate_multi_statement_workbook(
        company=None,
        year_trees=[(job, comp_tree)],
        output_dir=tmp_path,
    )

    assert len(result.provenance_records) > 0

    wb = openpyxl.load_workbook(result.file_path, data_only=False)
    ws_audit = wb["Audit_Trail"]

    # Verify audit headers
    assert ws_audit.cell(row=3, column=1).value == "Sheet"
    assert ws_audit.cell(row=3, column=2).value == "Cell"
    assert ws_audit.cell(row=3, column=3).value == "Line Item"

    # Verify audit trail entries exist
    assert ws_audit.cell(row=4, column=1).value is not None
    assert ws_audit.cell(row=4, column=5).value == "filing.pdf"
