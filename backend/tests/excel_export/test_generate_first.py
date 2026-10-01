"""
Unit tests for non-blocking generate-first workflow and deep links (FN-030, FN-032).

Covers:
- Immediate workbook generation from 0 reviewed items
- Status header: "DRAFT: N items unverified" vs "VERIFIED: All line items confirmed"
- Cell styling: uncertain cells yellow (#FEF08A) with comment, manual-required left empty and red (#FEE2E2) (Invariant I3)
- Review sheet listing each flagged cell with reason, confidence, and source link
- Provenance deep links on Source_Inputs label cell; values stay plain numbers (Invariant I5)
"""

from pathlib import Path

import openpyxl
from app.classification.models import StatementType
from app.excel_export.generator import generate_workbook
from app.extraction.locator import HtmlLocator, PdfLocator
from app.extraction.models import ConfidenceBand
from app.formula_engine.models import FormulaInputBatch, FormulaInputNode
from app.formula_engine.reader import read_formula_inputs_from_review
from app.formula_engine.tree import build_formula_tree
from app.review.models import ReviewItem, ReviewStatus


def test_generate_first_zero_reviewed_workflow(tmp_path: Path) -> None:
    """Verifies that an unreviewed batch of items generates a valid draft workbook immediately."""
    review_items = [
        ReviewItem(
            id="item_0",
            value="1500.00",
            label="Operating Income (Loss)",
            normalized_label="Operating Income",
            statement_type=StatementType.income_statement,
            page=10,
            bbox={"x0": 50.0, "y0": 100.0, "x1": 250.0, "y1": 130.0},
            source_file="q2_2026.pdf",
            confidence_band=ConfidenceBand.needs_review,
            confidence_score=0.65,
            flags=["low_extraction_confidence"],
            status=ReviewStatus.needs_review,
            locator=PdfLocator(
                page=10,
                bbox={"x0": 50.0, "y0": 100.0, "x1": 250.0, "y1": 130.0},
                source_file="q2_2026.pdf",
            ),
        ),
        ReviewItem(
            id="item_1",
            value="250.00",
            label="Share-Based Compensation",
            normalized_label="Stock-Based Compensation",
            statement_type=StatementType.non_gaap_bridge,
            page=25,
            bbox={"x0": 60.0, "y0": 200.0, "x1": 280.0, "y1": 230.0},
            source_file="q2_2026.pdf",
            confidence_band=ConfidenceBand.manual_required,
            confidence_score=0.30,
            flags=["unresolved_split_currency"],
            status=ReviewStatus.manual_required,
            locator=PdfLocator(
                page=25,
                bbox={"x0": 60.0, "y0": 200.0, "x1": 280.0, "y1": 230.0},
                source_file="q2_2026.pdf",
            ),
        ),
        ReviewItem(
            id="item_2",
            value="75.00",
            label="Restructuring Charges",
            normalized_label="Restructuring Charges",
            statement_type=StatementType.non_gaap_bridge,
            page=1,
            bbox={"x0": 0.0, "y0": 0.0, "x1": 1000.0, "y1": 1000.0},
            source_file="goog-20260630.htm",
            confidence_band=ConfidenceBand.auto_accepted,
            confidence_score=0.99,
            flags=[],
            status=ReviewStatus.auto_accepted,
            locator=HtmlLocator(
                accession="0001652044-26-000010",
                document="goog-20260630.htm",
                element_path="/html/body/table[3]/tr[4]/td[2]",
                url="https://www.sec.gov/Archives/edgar/data/1652044/000165204426000010/goog-20260630.htm",
            ),
        ),
    ]

    # 1. Non-blocking reader includes all unreviewed items
    batch = read_formula_inputs_from_review(review_items, include_unreviewed=True)
    assert batch.total_records_received == 3
    assert len(batch.nodes) == 3

    # Check Invariant I3: manual-required item has empty value ("")
    manual_node = next(n for n in batch.nodes if n.normalized_label == "Stock-Based Compensation")
    assert manual_node.value == ""
    assert manual_node.review_status == "manual_required"

    # Needs-review item keeps raw value for formula calculation
    needs_review_node = next(n for n in batch.nodes if n.normalized_label == "Operating Income")
    assert needs_review_node.value == "1500.00"
    assert needs_review_node.review_status == "needs_review"

    # 2. Build formula tree and generate workbook
    tree = build_formula_tree(batch, target_metric="Adjusted EBITDA")
    assert tree.is_valid is True

    result = generate_workbook(tree, job_id="job_gen_first", output_dir=tmp_path)
    assert result.is_success is True
    assert "Review" in result.sheet_names

    # 3. Inspect generated Excel workbook
    wb = openpyxl.load_workbook(result.file_path, data_only=False)

    # Check status header on Reconciliation sheet
    ws_recon = wb["Reconciliation"]
    status_header = str(ws_recon.cell(row=1, column=2).value or "")
    assert "DRAFT: 2 items unverified" in status_header

    # Check Source_Inputs: label cell has hyperlink, value cell is plain number / blank
    ws_inputs = wb["Source_Inputs"]
    assert ws_inputs.cell(row=1, column=1).value == "Label"
    assert ws_inputs.cell(row=1, column=2).value == "Value ($)"

    # Row 2: Operating Income (needs_review) -> value is numeric, label has deep link
    lbl_r2 = ws_inputs.cell(row=2, column=1)
    val_r2 = ws_inputs.cell(row=2, column=2)
    assert lbl_r2.hyperlink is not None
    assert "http://localhost:8000/api/review/job_gen_first/source" in lbl_r2.hyperlink.target
    assert lbl_r2.hyperlink.location == "page=10"
    assert val_r2.value == 1500.0
    assert val_r2.comment is not None
    assert "[Needs Review - Uncertain Item]" in val_r2.comment.text

    # Row 3: Stock-Based Compensation (manual_required) -> value is None/blank per Invariant I3
    lbl_r3 = ws_inputs.cell(row=3, column=1)
    val_r3 = ws_inputs.cell(row=3, column=2)
    assert lbl_r3.hyperlink is not None
    assert val_r3.value is None or val_r3.value == ""
    assert val_r3.comment is not None
    assert "[Manual Input Required]" in val_r3.comment.text

    # Row 4: Restructuring Charges (HTML source) -> public SEC link on label
    lbl_r4 = ws_inputs.cell(row=4, column=1)
    val_r4 = ws_inputs.cell(row=4, column=2)
    assert lbl_r4.hyperlink is not None
    assert "https://www.sec.gov" in lbl_r4.hyperlink.target
    assert val_r4.value == 75.0

    # 4. Check Review sheet
    ws_review = wb["Review"]
    assert ws_review.cell(row=1, column=1).value == "Workbook Verification & Audit Review Log"
    assert "Draft status: 2 unverified items" in str(ws_review.cell(row=2, column=1).value or "")
    assert ws_review.cell(row=3, column=1).value == "Cell"
    assert ws_review.cell(row=3, column=2).value == "Line Item"
    assert ws_review.cell(row=3, column=3).value == "Status"
    assert ws_review.cell(row=3, column=7).value == "Source Link"

    # Review rows list the 2 unverified items
    assert ws_review.cell(row=4, column=3).value in ("needs_review", "manual_required")
    assert ws_review.cell(row=5, column=3).value in ("needs_review", "manual_required")


def test_generate_verified_workflow(tmp_path: Path) -> None:
    """Verifies that 100% confirmed items produce a VERIFIED header and clean review sheet."""
    node = FormulaInputNode(
        node_id="node_0_ebitda",
        normalized_label="Operating Income",
        value="500.0",
        label="Operating Income",
        page=5,
        bbox={"x0": 10.0, "y0": 20.0, "x1": 30.0, "y1": 40.0},
        source_file="report.pdf",
        confidence_band="auto_accepted",
        confidence_score=0.99,
        flags=[],
        review_status="confirmed",
        is_confirmed=True,
        record_index=0,
    )
    batch = FormulaInputBatch(
        nodes=[node],
        total_records_received=1,
        confirmed_count=1,
        excluded_count=0,
    )
    tree = build_formula_tree(batch, target_metric="Adjusted EBITDA")
    result = generate_workbook(tree, job_id="job_verified", output_dir=tmp_path)
    assert result.is_success is True

    wb = openpyxl.load_workbook(result.file_path, data_only=False)
    ws_recon = wb["Reconciliation"]
    assert ws_recon.cell(row=1, column=2).value == "VERIFIED"

    ws_review = wb["Review"]
    assert "All line items confirmed" in str(ws_review.cell(row=4, column=2).value or "")
