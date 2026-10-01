"""
Unit tests for multi-period workbook default, restatements, gaps, and LTM formulas (FN-031).

Covers:
- Multi-period layout with quarterly columns (e.g. Q1-Q4) and LTM formula column (=SUM(B:E))
- Restatement detection: latest filing wins and cell comment records "Restated, was X"
- Missing periods/items left blank as explicit gaps (Invariant I3, never 0)
"""

from pathlib import Path

import openpyxl
from app.excel_export.multi_year_generator import generate_multi_year_workbook
from app.formula_engine.models import (
    FormulaInputBatch,
    FormulaInputNode,
    FormulaNode,
    FormulaNodeType,
    FormulaTree,
)
from app.ingestion.models import CompanyRecord, JobRecord, JobStatus


def _make_input_node(
    node_id: str, label: str, value: str, page: int = 1
) -> FormulaInputNode:
    return FormulaInputNode(
        node_id=node_id,
        label=label,
        normalized_label=label,
        value=value,
        page=page,
        bbox={"x0": 100.0, "y0": 200.0, "x1": 500.0, "y1": 250.0},
        source_file=f"filing_page_{page}.pdf",
        record_index=0,
    )


def _make_tree(
    target_metric: str, leaves_data: list[tuple[str, str, str]]
) -> FormulaTree:
    leaves: list[FormulaNode] = []
    for node_id, label, value in leaves_data:
        src = _make_input_node(node_id, label, value)
        leaves.append(
            FormulaNode(
                node_id=node_id,
                label=label,
                node_type=FormulaNodeType.leaf,
                source_node=src,
                children=[],
            )
        )

    root = FormulaNode(
        node_id=f"root-{target_metric}",
        label=target_metric,
        node_type=FormulaNodeType.calculated_root,
        children=leaves,
    )

    batch = FormulaInputBatch(
        nodes=[leaf.source_node for leaf in leaves if leaf.source_node is not None],
        total_records_received=len(leaves),
        confirmed_count=len(leaves),
        excluded_count=0,
    )

    return FormulaTree(
        root=root,
        leaves=leaves,
        target_metric=target_metric,
        batch=batch,
        is_valid=True,
    )


def test_multi_period_quarterly_with_ltm_and_restatement(tmp_path: Path) -> None:
    """Verifies that 4 quarterly periods generate an LTM formula column and restatements are flagged."""
    company = CompanyRecord(
        company_id="comp-goog",
        name="Alphabet Inc.",
        ticker="GOOGL",
        created_at="2026-08-23T12:00:00Z",
    )

    # Q1 2024 initial filing (Restructuring = 25.0)
    tree_q1_orig = _make_tree(
        "Adjusted EBITDA",
        [
            ("sbc-q1", "Stock-Based Compensation", "500.0"),
            ("rest-q1", "Restructuring Charges", "25.0"),
        ],
    )
    job_q1_orig = JobRecord(
        job_id="job-q1-orig",
        filename="goog_10q_q1.htm",
        file_size_bytes=1000,
        status=JobStatus.done,
        target_metric="Adjusted EBITDA",
        submitted_at="2024-05-01T12:00:00Z",
        filing_year=2024,
        period="Q1 2024",
    )

    # Q1 2024 restated in subsequent filing (Restructuring updated to 35.0, latest filing wins!)
    tree_q1_restated = _make_tree(
        "Adjusted EBITDA",
        [
            ("sbc-q1", "Stock-Based Compensation", "500.0"),
            ("rest-q1", "Restructuring Charges", "35.0"),
        ],
    )
    job_q1_restated = JobRecord(
        job_id="job-q1-restated",
        filename="goog_10k_2024.htm",
        file_size_bytes=2000,
        status=JobStatus.done,
        target_metric="Adjusted EBITDA",
        submitted_at="2025-02-01T12:00:00Z",
        filing_year=2024,
        period="Q1 2024",
    )

    # Q2 2024
    tree_q2 = _make_tree(
        "Adjusted EBITDA",
        [
            ("sbc-q2", "Stock-Based Compensation", "520.0"),
            ("rest-q2", "Restructuring Charges", "10.0"),
        ],
    )
    job_q2 = JobRecord(
        job_id="job-q2",
        filename="goog_10q_q2.htm",
        file_size_bytes=1000,
        status=JobStatus.done,
        target_metric="Adjusted EBITDA",
        submitted_at="2024-08-01T12:00:00Z",
        filing_year=2024,
        period="Q2 2024",
    )

    # Q3 2024
    tree_q3 = _make_tree(
        "Adjusted EBITDA",
        [
            ("sbc-q3", "Stock-Based Compensation", "540.0"),
            # Restructuring Charges absent in Q3 -> explicit gap, blank cell
        ],
    )
    job_q3 = JobRecord(
        job_id="job-q3",
        filename="goog_10q_q3.htm",
        file_size_bytes=1000,
        status=JobStatus.done,
        target_metric="Adjusted EBITDA",
        submitted_at="2024-11-01T12:00:00Z",
        filing_year=2024,
        period="Q3 2024",
    )

    # Q4 2024
    tree_q4 = _make_tree(
        "Adjusted EBITDA",
        [
            ("sbc-q4", "Stock-Based Compensation", "560.0"),
            ("rest-q4", "Restructuring Charges", "15.0"),
        ],
    )
    job_q4 = JobRecord(
        job_id="job-q4",
        filename="goog_10k_2024.htm",
        file_size_bytes=2000,
        status=JobStatus.done,
        target_metric="Adjusted EBITDA",
        submitted_at="2025-02-01T12:00:00Z",
        filing_year=2024,
        period="Q4 2024",
    )

    jobs = [
        (job_q1_orig, tree_q1_orig),
        (job_q1_restated, tree_q1_restated),
        (job_q2, tree_q2),
        (job_q3, tree_q3),
        (job_q4, tree_q4),
    ]

    result = generate_multi_year_workbook(company, jobs, output_dir=tmp_path)
    assert result.is_success is True

    wb = openpyxl.load_workbook(result.file_path, data_only=False)
    ws = wb["Multi_Year_Model"]

    # Header Row (Row 2): Line Item | Q1 2024 | Q2 2024 | Q3 2024 | Q4 2024 | LTM
    assert ws.cell(row=2, column=1).value == "Line Item"
    assert ws.cell(row=2, column=2).value == "Q1 2024"
    assert ws.cell(row=2, column=3).value == "Q2 2024"
    assert ws.cell(row=2, column=4).value == "Q3 2024"
    assert ws.cell(row=2, column=5).value == "Q4 2024"
    assert ws.cell(row=2, column=6).value == "LTM"

    # Line Item 1: Stock-Based Compensation
    assert ws.cell(row=3, column=1).value == "Stock-Based Compensation"
    assert ws.cell(row=3, column=2).value == 500.0
    assert ws.cell(row=3, column=3).value == 520.0
    assert ws.cell(row=3, column=4).value == 540.0
    assert ws.cell(row=3, column=5).value == 560.0
    # LTM formula: =SUM(B3:E3)
    assert str(ws.cell(row=3, column=6).value) == "=SUM(B3:E3)"

    # Line Item 2: Restructuring Charges
    assert ws.cell(row=4, column=1).value == "Restructuring Charges"
    # Latest filing won: 35.0 (not 25.0)
    assert ws.cell(row=4, column=2).value == 35.0
    # Restatement comment attached: "Restated, was 25.0"
    restated_comment = ws.cell(row=4, column=2).comment
    assert restated_comment is not None
    assert "Restated, was 25.0" in restated_comment.text

    # Q3 was absent: blank cell (not zero!) per Invariant I3
    assert ws.cell(row=4, column=4).value is None or ws.cell(row=4, column=4).value == ""

    # LTM formula for Restructuring Charges: =SUM(B4:E4)
    assert str(ws.cell(row=4, column=6).value) == "=SUM(B4:E4)"

    # Total row (Row 5): Total formulas
    assert str(ws.cell(row=5, column=2).value) == "=SUM(B3:B4)"
    assert str(ws.cell(row=5, column=6).value) == "=SUM(F3:F4)"
