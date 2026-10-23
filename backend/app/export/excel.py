"""
Excel export of one issuer (P4-07).

Sheets:
- ``Bridge``: every bridge line by quarter in USD millions, with the as-reported
  string and the sec.gov link in each cell's comment;
- ``Normalized``: reported headline, each excluded add-back, and normalized
  EBITDA as a live formula (reported - SUM(excluded));
- ``Sources``: release, accession, filing date and link per quarter.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

_HEAD = Font(bold=True, color="FFFFFF")
_FILL = PatternFill("solid", fgColor="1F2A44")
_NUM = '#,##0.0;(#,##0.0);"-"'
_MULT = {"units": 1.0, "thousands": 1e3, "millions": 1e6, "billions": 1e9}


def _m(v: float | None) -> float | None:
    return None if v is None else v / 1_000_000


def issuer_workbook(doc: dict[str, Any], policy_version: str) -> bytes:
    wb = Workbook()
    quarters = doc["quarters"]

    ws = wb.active
    if ws is None:
        raise RuntimeError("openpyxl workbook has no active sheet")
    ws.title = "Bridge"
    ws.append(
        [f"{doc['ticker']} - {doc['name']}: reported EBITDA bridges (USD millions)"]
    )
    ws.append(["Line", "Role", *[q["end_date"] for q in quarters]])
    for c in ws[2]:
        c.font, c.fill = _HEAD, _FILL
    order: list[tuple[str, str]] = []
    for q in reversed(quarters):
        for ln in q["lines"]:
            key = (ln["label"], ln["role"])
            if key not in order:
                order.append(key)
    for label, role in order:
        ws.append([label, role])
        r = ws.max_row
        for j, q in enumerate(quarters):
            hit = next(
                (x for x in q["lines"] if x["label"] == label and x["role"] == role),
                None,
            )
            if hit is None:
                continue
            cell = ws.cell(row=r, column=3 + j, value=_m(hit["value_usd"]))
            cell.number_format = _NUM
            cell.comment = Comment(
                f"As reported: {hit['value_raw']} ({q['scale']})\n{hit['url'] or ''}",
                "Footnote",
            )
    ws.column_dimensions["A"].width = 60

    ns = wb.create_sheet("Normalized")
    ns.append([f"Normalized EBITDA, policy {policy_version} (preview; USD millions)"])
    ns.append(["Item", *[q["end_date"] for q in quarters]])
    for c in ns[2]:
        c.font, c.fill = _HEAD, _FILL
    ns.append(["Reported headline", *[_m(q["reported_usd"]) for q in quarters]])
    excluded: list[str] = []
    for q in quarters:
        for ln in q["normalization"]["lines"]:
            if ln["treatment"] == "exclude" and ln["label"] not in excluded:
                excluded.append(ln["label"])
    first_excl = ns.max_row + 1
    for label in excluded:
        row: list[Any] = [f"Less: {label}"]
        for q in quarters:
            hit = next(
                (
                    x
                    for x in q["normalization"]["lines"]
                    if x["label"] == label and x["treatment"] == "exclude"
                ),
                None,
            )
            if hit is None or hit["contribution"] is None or q["scale"] not in _MULT:
                row.append(None)
            else:
                row.append(_m(float(hit["contribution"]) * _MULT[q["scale"]]))
        ns.append(row)
    last_excl = ns.max_row
    ns.append(["Normalized EBITDA"])
    out_row = ns.max_row
    for j, q in enumerate(quarters):
        col = get_column_letter(2 + j)
        norm = q["normalization"]
        if norm["status"] != "computed":
            ns.cell(row=out_row, column=2 + j, value=f"not computed: {norm['reason']}")
            continue
        formula = f"={col}3"
        if excluded:
            formula += f"-SUM({col}{first_excl}:{col}{last_excl})"
        ns.cell(row=out_row, column=2 + j, value=formula)
    for cells in ns.iter_rows(min_row=3, min_col=2):
        for c in cells:
            c.number_format = _NUM
    ns.cell(row=out_row, column=1).font = Font(bold=True)
    ns.column_dimensions["A"].width = 60

    ss = wb.create_sheet("Sources")
    ss.append(["Quarter end", "Accession", "Filed", "Headline label", "Source"])
    for c in ss[1]:
        c.font, c.fill = _HEAD, _FILL
    for q in quarters:
        ss.append(
            [
                q["end_date"],
                q["accession"],
                q["filed"],
                q["metric_label"],
                q["exhibit_url"],
            ]
        )
    ss.column_dimensions["E"].width = 100

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
